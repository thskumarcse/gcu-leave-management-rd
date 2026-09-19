"""
Business logic for importing the university employee master CSV/Excel.
"""
import csv
import datetime
import shutil
import tempfile
import time
from pathlib import Path
import re
from dataclasses import dataclass, field

from django.core.exceptions import ValidationError
from openpyxl import load_workbook

from .models import Employee

REQUIRED_COLUMNS = [
    'emp_id',
    'Name',
    'Department',
    'Group Name',
    'User Type',
    'Designation',
    'Joining Date',
    'Email ID',
    'Mobile No.',
    'Date of Birth',
    'Sub Department',
    'Academy',
]

MINIMUM_COLUMNS = ['emp_id', 'Name']

SYSTEM_EMP_IDS = {
    'SYSADMIN',
    'APPLICANT_PORTAL_USER',
    'mis',
    'Admin',
    'Management',
    'CAP123',
    'T100',
    'SeroCST',
    'sambit.kundu',
    'aman.sharma',
    'serosunny',
    'serorakesh',
    'AI Admin',
}

PLACEHOLDER_NAME_MARKERS = (
    'APP_FIRST_NAME',
    'DEMO FILE',
    'TEST ACCOUNT',
    'ERP ADMIN',
)

GROUP_TO_DESIGNATION_TYPE = {
    'faculty': Employee.DesignationType.FACULTY,
    'teaching': Employee.DesignationType.FACULTY,
    'admin': Employee.DesignationType.STAFF,
    'administrative': Employee.DesignationType.STAFF,
    'other employee': Employee.DesignationType.STAFF,
}

GROUP_NAME_LABEL = {
    'faculty': 'Faculty',
    'teaching': 'Faculty',
    'admin': 'Admin',
    'administrative': 'Admin',
    'other employee': 'Other Employee',
}

HEADER_ALIASES = {
    'emp_id': 'emp_id',
    'employee id': 'emp_id',
    'employeeid': 'emp_id',
    'name': 'Name',
    'email': 'Email ID',
    'email id': 'Email ID',
    'group name': 'Group Name',
    'type': 'Group Name',
    'user type': 'User Type',
    'designation': 'Designation',
    'joining date': 'Joining Date',
    'mobile no.': 'Mobile No.',
    'mobile no': 'Mobile No.',
    'mobile': 'Mobile No.',
    'date of birth': 'Date of Birth',
    'dob': 'Date of Birth',
    'sub department': 'Sub Department',
    'academy': 'Academy',
    'status': 'Status',
    'department': 'Department',
    'school': 'Department',
}

SKIP_HEADERS = {
    'sl',
    's.l',
    's.no',
    's.no.',
    'sno',
    'serial no',
    'serial no.',
    'campus',
}

DATE_FORMATS = ('%d-%m-%Y', '%Y-%m-%d', '%d/%m/%Y')
BLANK_TOKENS = {
    '', '-', '--', 'NA', 'N/A', '.', 'null', 'none',
    '—', '–', '\u2014', '\u2013',
}
SPREADSHEET_SUFFIXES = {'.xlsx', '.xlsm'}
CREATE_DEFAULTS = {
    'dob': None,
    'designation': '',
    'designation_type': Employee.DesignationType.STAFF,
    'group_name': '',
    'user_type': '',
    'department': '',
    'sub_department': '',
    'academy': '',
    'email': '',
    'mobile': '',
    'joining_date': None,
    'status': Employee.Status.ACTIVE,
}


@dataclass
class RowIssue:
    row_number: int
    emp_id: str
    reason: str


@dataclass
class ImportReport:
    imported: int = 0
    updated: int = 0
    skipped: list = field(default_factory=list)
    deactivated: int = 0
    total_rows: int = 0

    @property
    def skipped_count(self):
        return len(self.skipped)


class CsvFormatError(Exception):
    """Raised when the CSV itself is unusable (missing file/columns)."""


def _cell_to_str(value):
    if value is None:
        return ''
    if isinstance(value, datetime.datetime):
        return value.strftime('%d-%m-%Y')
    if isinstance(value, datetime.date):
        return value.strftime('%d-%m-%Y')
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)
    return str(value).strip()


def _blankish(value):
    value = _cell_to_str(value)
    if value.lower() in BLANK_TOKENS or value in BLANK_TOKENS:
        return ''
    return value


def _clean_name(value):
    return ' '.join(_blankish(value).split())


def _parse_date(value, field_name):
    value = _blankish(value)
    if not value:
        return None
    for fmt in DATE_FORMATS:
        try:
            return datetime.datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    raise ValidationError(f'{field_name} "{value}" is not a valid date.')


def _parse_mobile(value):
    value = _blankish(value)
    if not value:
        return ''
    if re.search(r'[eE]', value):
        try:
            value = str(int(float(value)))
        except ValueError:
            return ''
    if value.startswith('-'):
        value = value[1:]
    digits = re.sub(r'\D', '', value)
    if len(digits) not in (10, 11, 12, 13):
        return ''
    if len(set(digits)) == 1:
        return ''
    if digits in {'1111111020', '0000000000'}:
        return ''
    return digits


def _parse_email(value):
    value = _blankish(value).lower()
    if not value or '@' not in value:
        return ''
    return value


def _parse_status(value):
    value = _blankish(value)
    if not value:
        return Employee.Status.ACTIVE
    key = value.lower()
    if key == 'active':
        return Employee.Status.ACTIVE
    if key == 'inactive':
        return Employee.Status.INACTIVE
    raise ValidationError(f'Status "{value}" is not Active or Inactive.')


def _map_designation_type(group_name):
    mapped = GROUP_TO_DESIGNATION_TYPE.get(group_name.lower())
    if mapped:
        return mapped
    return Employee.DesignationType.STAFF


def _normalize_group_name(group_name):
    raw = _blankish(group_name)
    if not raw:
        return ''
    return GROUP_NAME_LABEL.get(raw.lower(), raw)


def _is_system_row(emp_id, name):
    if emp_id in SYSTEM_EMP_IDS:
        return True
    upper_name = name.upper()
    return any(marker.upper() in upper_name for marker in PLACEHOLDER_NAME_MARKERS)


def _canonicalize_headers(raw_headers):
    """Map file headers onto the employee.csv column names."""
    cleaned = [_cell_to_str(header) for header in (raw_headers or [])]
    lowered = [header.lower() for header in cleaned]
    has_school = 'school' in lowered
    canonical = []
    seen = set()
    for header in cleaned:
        key = header.lower()
        if not key or key in SKIP_HEADERS:
            canonical.append(None)
            continue
        if has_school and key == 'school':
            name = 'Department'
        elif has_school and key == 'department':
            name = 'Sub Department'
        else:
            name = HEADER_ALIASES.get(key, header)
        if name in seen:
            canonical.append(None)
            continue
        seen.add(name)
        canonical.append(name)
    return canonical


def _has_column(fieldnames, column):
    return column in (fieldnames or [])


def _validate_row(row, fieldnames=None):
    fieldnames = set(fieldnames or row.keys())
    emp_id = _blankish(row.get('emp_id'))
    name = _clean_name(row.get('Name'))

    if not emp_id:
        raise ValidationError('Missing required value: emp_id.')
    if not name:
        raise ValidationError('Missing required value: Name.')
    if _is_system_row(emp_id, name):
        raise ValidationError('System / ERP account — not imported as a university employee.')

    cleaned = {
        'emp_id': emp_id,
        'name': name,
    }

    if _has_column(fieldnames, 'Group Name'):
        raw_group = _blankish(row.get('Group Name'))
        cleaned['group_name'] = _normalize_group_name(raw_group)
        cleaned['designation_type'] = _map_designation_type(raw_group)

    if _has_column(fieldnames, 'User Type'):
        cleaned['user_type'] = _blankish(row.get('User Type'))

    if _has_column(fieldnames, 'Designation'):
        cleaned['designation'] = _blankish(row.get('Designation'))

    if _has_column(fieldnames, 'Department'):
        cleaned['department'] = _blankish(row.get('Department'))

    if _has_column(fieldnames, 'Sub Department'):
        cleaned['sub_department'] = _blankish(row.get('Sub Department'))

    if _has_column(fieldnames, 'Academy'):
        cleaned['academy'] = _blankish(row.get('Academy'))

    if _has_column(fieldnames, 'Email ID'):
        cleaned['email'] = _parse_email(row.get('Email ID'))

    if _has_column(fieldnames, 'Mobile No.'):
        cleaned['mobile'] = _parse_mobile(row.get('Mobile No.'))

    if _has_column(fieldnames, 'Date of Birth'):
        cleaned['dob'] = _parse_date(row.get('Date of Birth'), 'Date of Birth')

    if _has_column(fieldnames, 'Joining Date'):
        cleaned['joining_date'] = _parse_date(row.get('Joining Date'), 'Joining Date')

    if _has_column(fieldnames, 'Status'):
        cleaned['status'] = _parse_status(row.get('Status'))
    else:
        # Original employee.csv has no Status column and always marks rows Active.
        cleaned['status'] = Employee.Status.ACTIVE

    return cleaned


def _row_values(canonical_headers, values):
    row = {}
    for header, value in zip(canonical_headers, values):
        if not header:
            continue
        row[header] = _cell_to_str(value)
    return row


def _is_blank_row(values):
    return not any(_cell_to_str(value) for value in values)


def _required_columns_for(fieldnames):
    present = set(fieldnames or [])
    if present.issuperset(REQUIRED_COLUMNS):
        return REQUIRED_COLUMNS
    return MINIMUM_COLUMNS


def _load_xlsx_workbook(file_path):
    last_error = None
    tmp_path = None
    for _attempt in range(4):
        try:
            return load_workbook(file_path, read_only=True, data_only=True)
        except FileNotFoundError:
            raise CsvFormatError(f'Spreadsheet not found at {file_path}.')
        except PermissionError as exc:
            last_error = exc
            handle = tempfile.NamedTemporaryFile(suffix=Path(file_path).suffix, delete=False)
            handle.close()
            tmp_path = handle.name
            try:
                shutil.copyfile(file_path, tmp_path)
                return load_workbook(tmp_path, read_only=True, data_only=True)
            except OSError as copy_exc:
                last_error = copy_exc
                time.sleep(1)
        except Exception as exc:
            last_error = exc
            break
    raise CsvFormatError(f'Could not read spreadsheet at {file_path}: {last_error}')


def _read_spreadsheet_rows(file_path):
    workbook = _load_xlsx_workbook(file_path)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(iterator)
        except StopIteration:
            raise CsvFormatError(f'Spreadsheet at {file_path} has no header row.')
        canonical = _canonicalize_headers(raw_headers)
        fieldnames = [name for name in canonical if name]
        required = _required_columns_for(fieldnames)
        missing = [column for column in required if column not in fieldnames]
        if missing:
            raise CsvFormatError(
                f'Spreadsheet is missing required column(s): {", ".join(missing)}.'
            )
        rows = []
        for offset, values in enumerate(iterator):
            if _is_blank_row(values):
                continue
            rows.append((offset + 2, _row_values(canonical, values)))
        return fieldnames, rows
    finally:
        workbook.close()


def _read_csv_rows(file_path):
    try:
        handle = open(file_path, newline='', encoding='utf-8-sig')
    except FileNotFoundError:
        raise CsvFormatError(f'CSV file not found at {file_path}.')

    with handle:
        reader = csv.reader(handle)
        try:
            raw_headers = next(reader)
        except StopIteration:
            raise CsvFormatError(f'CSV file at {file_path} has no header row.')
        canonical = _canonicalize_headers(raw_headers)
        fieldnames = [name for name in canonical if name]
        required = _required_columns_for(fieldnames)
        missing = [column for column in required if column not in fieldnames]
        if missing:
            raise CsvFormatError(
                f'CSV is missing required column(s): {", ".join(missing)}.'
            )
        rows = []
        for offset, values in enumerate(reader):
            if _is_blank_row(values):
                continue
            rows.append((offset + 2, _row_values(canonical, values)))
        return fieldnames, rows


def _save_employee(cleaned):
    existing = Employee.objects.filter(emp_id=cleaned['emp_id']).first()
    if existing is None:
        payload = {**CREATE_DEFAULTS, **cleaned}
        employee = Employee(**payload)
        employee.full_clean(exclude=['id'], validate_unique=False)
        employee.save()
        return True
    for field_name, value in cleaned.items():
        if field_name == 'emp_id':
            continue
        setattr(existing, field_name, value)
    existing.full_clean(exclude=['id'], validate_unique=False)
    existing.save()
    return False


def _import_employee_rows(fieldnames, rows, deactivate_missing=False):
    report = ImportReport()
    seen_emp_ids = {}

    for row_number, row in rows:
        report.total_rows += 1
        raw_emp_id = _blankish(row.get('emp_id'))

        if raw_emp_id and raw_emp_id in seen_emp_ids:
            report.skipped.append(RowIssue(
                row_number=row_number,
                emp_id=raw_emp_id,
                reason=(
                    f'Duplicate emp_id in CSV — already processed at '
                    f'row {seen_emp_ids[raw_emp_id]}.'
                ),
            ))
            continue

        try:
            cleaned = _validate_row(row, fieldnames)
        except ValidationError as exc:
            report.skipped.append(RowIssue(
                row_number=row_number,
                emp_id=raw_emp_id or '(missing)',
                reason='; '.join(exc.messages) if hasattr(exc, 'messages') else str(exc),
            ))
            continue

        seen_emp_ids[cleaned['emp_id']] = row_number

        try:
            created = _save_employee(cleaned)
        except ValidationError as exc:
            report.skipped.append(RowIssue(
                row_number=row_number,
                emp_id=cleaned['emp_id'],
                reason='; '.join(
                    f'{field}: {", ".join(msgs)}'
                    for field, msgs in exc.message_dict.items()
                ) if hasattr(exc, 'message_dict') else str(exc),
            ))
            continue

        if created:
            report.imported += 1
        else:
            report.updated += 1

    if deactivate_missing and seen_emp_ids:
        report.deactivated = Employee.objects.exclude(
            emp_id__in=seen_emp_ids.keys()
        ).exclude(
            status=Employee.Status.INACTIVE
        ).update(status=Employee.Status.INACTIVE)

    return report


def import_employees_from_csv(file_path, deactivate_missing=False):
    """
    Reads and imports the employee CSV at file_path. Per-row problems are
    collected as skipped rows. Only a structural problem with the CSV
    (missing file/columns) raises CsvFormatError.
    """
    fieldnames, rows = _read_csv_rows(file_path)
    return _import_employee_rows(fieldnames, rows, deactivate_missing=deactivate_missing)


def import_employees_from_file(file_path, deactivate_missing=False):
    """Import employees from CSV or Excel, upserting on emp_id."""
    suffix = Path(file_path).suffix.lower()
    if suffix in SPREADSHEET_SUFFIXES:
        fieldnames, rows = _read_spreadsheet_rows(file_path)
        return _import_employee_rows(fieldnames, rows, deactivate_missing=deactivate_missing)
    return import_employees_from_csv(file_path, deactivate_missing=deactivate_missing)
