from datetime import date, timedelta
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.test import SimpleTestCase
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from apps.accounts.models import Account
from apps.accounts.services import provision_account
from apps.approvals.models import ApprovalAction, ApproverAssignment, EmployeeApprover
from apps.approvals.services import find_approver1, find_approver2
from apps.employees.models import Employee
from apps.leaves.models import LeaveApplication, LeaveType
from apps.notifications.models import Notification


def make_employee(**overrides):
    defaults = {
        'emp_id': 'GCU020041',
        'name': 'Bhabajit Baruah',
        'dob': '1988-05-06',
        'designation': 'Assistant Professor',
        'designation_type': Employee.DesignationType.FACULTY,
        'group_name': 'Faculty',
        'department': 'School Of Engineering and Technology',
        'sub_department': 'Department of Mechanical Engineering',
        'email': 'bhabajit_me@gcuniversity.ac.in',
        'mobile': '7002128161',
        'joining_date': '2012-08-01',
        'status': Employee.Status.ACTIVE,
    }
    defaults.update(overrides)
    return Employee.objects.create(**defaults)


def login_as(client, emp_id, password=None):
    password = password or settings.DEFAULT_EMPLOYEE_PASSWORD
    login = client.post(
        '/api/v1/auth/login/',
        {'emp_id': emp_id, 'password': password},
        format='json',
    )
    if login.status_code != 200 and password == settings.DEFAULT_EMPLOYEE_PASSWORD:
        return login_as(client, emp_id, 'NewPass#2026')
    client.credentials(HTTP_AUTHORIZATION=f'Bearer {login.data["data"]["access"]}')
    if login.data['data']['user']['must_change_password']:
        client.post(
            '/api/v1/auth/change-password/',
            {
                'current_password': password,
                'new_password': 'NewPass#2026',
                'confirm_password': 'NewPass#2026',
            },
            format='json',
        )
        password = 'NewPass#2026'
        login = client.post(
            '/api/v1/auth/login/',
            {'emp_id': emp_id, 'password': password},
            format='json',
        )
        client.credentials(HTTP_AUTHORIZATION=f'Bearer {login.data["data"]["access"]}')
    return password


class NotificationAndDashboardTests(APITestCase):
    def setUp(self):
        self.faculty = make_employee()
        self.hod = make_employee(
            emp_id='GCU020014',
            name='Debarshi Mallick',
            designation='Associate Professor & HOD',
            email='debarshi_me@gcuniversity.ac.in',
            mobile='7002474854',
        )
        self.vc = make_employee(
            emp_id='GCU090001',
            name='Vice Chancellor',
            designation='Vice-Chancellor',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='',
            email='vc@gcuniversity.ac.in',
            mobile='9000000001',
        )
        ApproverAssignment.objects.create(
            role=ApproverAssignment.Role.HOD,
            department=self.faculty.department,
            sub_department=self.faculty.sub_department,
            employee=self.hod,
        )
        ApproverAssignment.objects.create(
            role=ApproverAssignment.Role.VC,
            department='',
            sub_department='',
            employee=self.vc,
        )
        self.casual, _ = LeaveType.objects.get_or_create(
            code='CL',
            defaults={
                'name': 'Casual Leave',
                'max_days_per_year': 12,
                'applicable_to': LeaveType.ApplicableTo.ALL,
                'sort_order': 1,
            },
        )
        self.start = date.today() + timedelta(days=21)

    def test_apply_notifies_hod_and_decision_notifies_applicant(self):
        login_as(self.client, 'GCU020041')
        created = self.client.post(
            '/api/v1/leaves/',
            {
                'leave_type_id': self.casual.id,
                'start_date': self.start.isoformat(),
                'end_date': self.start.isoformat(),
                'reason': 'Personal work in Guwahati.',
            },
            format='json',
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED)
        self.assertTrue(
            Notification.objects.filter(employee=self.hod, is_read=False).exists()
        )

        login_as(self.client, 'GCU020014')
        inbox = self.client.get('/api/v1/notifications/')
        self.assertGreaterEqual(inbox.data['data']['unread'], 1)
        leave_id = created.data['data']['id']
        self.client.post(
            f'/api/v1/approvals/{leave_id}/decide/',
            {'decision': 'REJECTED', 'remarks': 'Coverage is too thin.'},
            format='json',
        )
        self.assertTrue(
            Notification.objects.filter(employee=self.faculty, title__icontains='rejected').exists()
        )

        login_as(self.client, 'GCU020041')
        listed = self.client.get('/api/v1/notifications/')
        first_id = listed.data['data']['results'][0]['id']
        read = self.client.post(f'/api/v1/notifications/{first_id}/read/')
        self.assertEqual(read.status_code, status.HTTP_200_OK)
        self.assertEqual(read.data['data']['updated'], 1)

    def test_approver_dashboard_requires_role(self):
        login_as(self.client, 'GCU020041')
        denied = self.client.get('/api/v1/dashboard/')
        self.assertEqual(denied.status_code, status.HTTP_403_FORBIDDEN)

        login_as(self.client, 'GCU020014')
        allowed = self.client.get('/api/v1/dashboard/')
        self.assertEqual(allowed.status_code, status.HTTP_200_OK)
        self.assertIn('pending', allowed.data['data'])
        self.assertTrue(allowed.data['data']['is_hod'])


class AdminApiTests(APITestCase):
    def setUp(self):
        self.faculty = make_employee()
        self.admin = make_employee(
            emp_id='GCU090001',
            name='Vice Chancellor',
            designation='Vice-Chancellor',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='',
            email='vc@gcuniversity.ac.in',
            mobile='9000000001',
        )
        login_as(self.client, 'GCU090001')
        account = Account.objects.get(employee=self.admin)
        account.is_admin = True
        account.save(update_fields=['is_admin', 'updated_at'])
        account.user.is_staff = True
        account.user.save(update_fields=['is_staff'])
        login_as(self.client, 'GCU090001', 'NewPass#2026')

    def test_non_admin_cannot_access_admin_api(self):
        login_as(self.client, 'GCU020041')
        response = self.client.get('/api/v1/admin/overview/')
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_admin_can_manage_leave_types_and_employees(self):
        overview = self.client.get('/api/v1/admin/overview/')
        self.assertEqual(overview.status_code, status.HTTP_200_OK)
        self.assertGreaterEqual(overview.data['data']['employees'], 2)

        created = self.client.post(
            '/api/v1/admin/leave-types/',
            {
                'name': 'Study Leave',
                'code': 'STL',
                'applicable_to': 'ALL',
                'max_days_per_year': 10,
                'sort_order': 20,
            },
            format='json',
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED)

        patched = self.client.patch(
            '/api/v1/admin/access/GCU020041/',
            {'is_admin': True},
            format='json',
        )
        self.assertEqual(patched.status_code, status.HTTP_200_OK)
        self.assertTrue(patched.data['data']['is_admin'])

        status_patch = self.client.patch(
            '/api/v1/admin/employees/GCU020041/',
            {'status': 'Active'},
            format='json',
        )
        self.assertEqual(status_patch.status_code, status.HTTP_200_OK)

        assignment = self.client.post(
            '/api/v1/admin/approvers/',
            {
                'role': 'HOD',
                'emp_id': 'GCU020041',
                'department': self.faculty.department,
                'sub_department': self.faculty.sub_department,
            },
            format='json',
        )
        self.assertIn(assignment.status_code, (status.HTTP_200_OK, status.HTTP_201_CREATED))

        report = self.client.get('/api/v1/admin/reports/')
        self.assertEqual(report.status_code, status.HTTP_200_OK)
        audit = self.client.get('/api/v1/admin/audit/')
        self.assertEqual(audit.status_code, status.HTTP_200_OK)

    def test_admin_can_patch_employee_master_fields(self):
        patched = self.client.patch(
            '/api/v1/admin/employees/GCU020041/',
            {
                'name': 'Bhabajit Baruah Corrected',
                'designation': 'Associate Professor',
                'sub_department': 'Department of Civil Engineering',
                'email': 'bhabajit.civil@gcuniversity.ac.in',
                'status': 'Active',
                'designation_type': 'FACULTY',
                'emp_id': 'GCU999999',
            },
            format='json',
        )
        self.assertEqual(patched.status_code, status.HTTP_200_OK)
        self.assertTrue(patched.data['success'])
        data = patched.data['data']
        self.assertEqual(data['name'], 'Bhabajit Baruah Corrected')
        self.assertEqual(data['designation'], 'Associate Professor')
        self.assertEqual(data['sub_department'], 'Department of Civil Engineering')
        self.assertEqual(data['email'], 'bhabajit.civil@gcuniversity.ac.in')
        self.assertEqual(data['emp_id'], 'GCU020041')
        self.faculty.refresh_from_db()
        self.assertEqual(self.faculty.emp_id, 'GCU020041')
        self.assertEqual(self.faculty.name, 'Bhabajit Baruah Corrected')
        self.assertEqual(self.faculty.sub_department, 'Department of Civil Engineering')

        login_as(self.client, 'GCU020041')
        denied = self.client.patch(
            '/api/v1/admin/employees/GCU020041/',
            {'name': 'Hacked'},
            format='json',
        )
        self.assertEqual(denied.status_code, status.HTTP_403_FORBIDDEN)

    def test_admin_can_add_employee(self):
        created = self.client.post(
            '/api/v1/admin/employees/',
            {
                'emp_id': 'GCU020198',
                'name': 'New Staff Member',
                'designation': 'Office Assistant',
                'department': 'Registrar Office',
            },
            format='json',
        )
        self.assertEqual(created.status_code, status.HTTP_201_CREATED)
        data = created.data['data']
        self.assertEqual(data['emp_id'], 'GCU020198')
        self.assertEqual(data['name'], 'New Staff Member')
        self.assertEqual(data['designation'], 'Office Assistant')
        self.assertEqual(data['sub_department'], 'Registrar Office')
        self.assertEqual(data['designation_type'], Employee.DesignationType.STAFF)
        self.assertTrue(Employee.objects.filter(emp_id='GCU020198').exists())

        duplicate = self.client.post(
            '/api/v1/admin/employees/',
            {
                'emp_id': 'gcu020198',
                'name': 'Someone Else',
            },
            format='json',
        )
        self.assertEqual(duplicate.status_code, status.HTTP_400_BAD_REQUEST)

        faculty = self.client.post(
            '/api/v1/admin/employees/',
            {
                'emp_id': 'GCU020199',
                'name': 'New Faculty Member',
                'designation': 'Assistant Professor',
                'department': 'Department of Mechanical Engineering',
            },
            format='json',
        )
        self.assertEqual(faculty.status_code, status.HTTP_201_CREATED)
        self.assertEqual(faculty.data['data']['designation_type'], Employee.DesignationType.FACULTY)

        login_as(self.client, 'GCU020041')
        denied = self.client.post(
            '/api/v1/admin/employees/',
            {'emp_id': 'GCU020200', 'name': 'Hacked'},
            format='json',
        )
        self.assertEqual(denied.status_code, status.HTTP_403_FORBIDDEN)

    def test_employee_filters_and_designation_query(self):
        listed = self.client.get(
            '/api/v1/admin/employees/',
            {'designation': 'Assistant Professor', 'status': 'Active'},
        )
        self.assertEqual(listed.status_code, status.HTTP_200_OK)
        emp_ids = [row['emp_id'] for row in listed.data['data']['results']]
        self.assertIn('GCU020041', emp_ids)

        options = self.client.get('/api/v1/admin/employees/filters/')
        self.assertEqual(options.status_code, status.HTTP_200_OK)
        self.assertIn(self.faculty.sub_department, options.data['data']['departments'])
        self.assertNotIn(self.faculty.department, options.data['data']['departments'])
        self.assertIn('Assistant Professor', options.data['data']['designations'])

        by_department = self.client.get(
            '/api/v1/admin/employees/',
            {
                'department': self.faculty.sub_department,
                'status': 'Active',
                'page': 1,
                'page_size': 500,
            },
        )
        self.assertEqual(by_department.status_code, status.HTTP_200_OK)
        dept_ids = [row['emp_id'] for row in by_department.data['data']['results']]
        self.assertIn('GCU020041', dept_ids)
        self.assertGreaterEqual(by_department.data['data']['count'], 1)

        by_sub = self.client.get(
            '/api/v1/admin/employees/',
            {
                'sub_department': self.faculty.sub_department,
                'status': 'Active',
                'page': 1,
                'page_size': 500,
            },
        )
        self.assertIn('GCU020041', [row['emp_id'] for row in by_sub.data['data']['results']])

        school_filter = self.client.get(
            '/api/v1/admin/employees/',
            {
                'department': self.faculty.department,
                'status': 'Active',
                'page': 1,
                'page_size': 500,
            },
        )
        school_ids = [row['emp_id'] for row in school_filter.data['data']['results']]
        self.assertNotIn('GCU020041', school_ids)

    def test_employee_blank_department_and_designation_filters(self):
        blank_dept = make_employee(
            emp_id='GCU010201',
            name='Blank Department Person',
            designation='Clerk',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='',
            email='blank_dept@gcuniversity.ac.in',
            mobile='7002000201',
        )
        blank_desig = make_employee(
            emp_id='GCU010202',
            name='Blank Designation Person',
            designation='',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='Establishment Section',
            email='blank_desig@gcuniversity.ac.in',
            mobile='7002000202',
        )

        by_blank_dept = self.client.get(
            '/api/v1/admin/employees/',
            {'sub_department': '__blank__', 'page_size': 500},
        )
        self.assertEqual(by_blank_dept.status_code, status.HTTP_200_OK)
        dept_ids = [row['emp_id'] for row in by_blank_dept.data['data']['results']]
        self.assertIn(blank_dept.emp_id, dept_ids)
        self.assertIn(self.admin.emp_id, dept_ids)
        self.assertNotIn(self.faculty.emp_id, dept_ids)

        by_flag = self.client.get(
            '/api/v1/admin/employees/',
            {'blank_department': '1', 'page_size': 500},
        )
        flag_ids = [row['emp_id'] for row in by_flag.data['data']['results']]
        self.assertIn(blank_dept.emp_id, flag_ids)
        self.assertNotIn(self.faculty.emp_id, flag_ids)

        by_blank_desig = self.client.get(
            '/api/v1/admin/employees/',
            {'designation': '__blank__', 'page_size': 500},
        )
        self.assertEqual(by_blank_desig.status_code, status.HTTP_200_OK)
        desig_ids = [row['emp_id'] for row in by_blank_desig.data['data']['results']]
        self.assertIn(blank_desig.emp_id, desig_ids)
        self.assertNotIn(self.faculty.emp_id, desig_ids)
        self.assertNotIn(blank_dept.emp_id, desig_ids)

        combined = self.client.get(
            '/api/v1/admin/employees/',
            {'designation': '__blank__', 'q': 'Blank Designation', 'page_size': 500},
        )
        combined_ids = [row['emp_id'] for row in combined.data['data']['results']]
        self.assertEqual(combined_ids, [blank_desig.emp_id])

    def test_admin_can_delete_employee_without_account_or_leaves(self):
        orphan = make_employee(
            emp_id='GCU010301',
            name='Orphan Record',
            designation='Assistant',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='Establishment Section',
            email='orphan@gcuniversity.ac.in',
            mobile='7002000301',
        )
        response = self.client.delete(f'/api/v1/admin/employees/{orphan.emp_id}/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertTrue(response.data['success'])
        self.assertEqual(response.data['data']['emp_id'], orphan.emp_id)
        self.assertFalse(Employee.objects.filter(emp_id=orphan.emp_id).exists())

    def test_admin_deletes_login_account_then_employee_when_safe(self):
        target = make_employee(
            emp_id='GCU010302',
            name='Has Login No Leaves',
            designation='Assistant',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='Establishment Section',
            email='haslogin@gcuniversity.ac.in',
            mobile='7002000302',
        )
        provision_account(target)
        self.assertTrue(Account.objects.filter(employee=target).exists())

        response = self.client.delete(f'/api/v1/admin/employees/{target.emp_id}/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertFalse(Employee.objects.filter(emp_id=target.emp_id).exists())
        self.assertFalse(Account.objects.filter(employee_id=target.id).exists())

    def test_cannot_delete_employee_with_leave_records(self):
        leave_type, _ = LeaveType.objects.get_or_create(
            code='CL',
            defaults={
                'name': 'Casual Leave',
                'max_days_per_year': 12,
                'applicable_to': LeaveType.ApplicableTo.ALL,
                'sort_order': 1,
            },
        )
        LeaveApplication.objects.create(
            employee=self.faculty,
            leave_type=leave_type,
            start_date=date.today() + timedelta(days=10),
            end_date=date.today() + timedelta(days=10),
            days=1,
            reason='Personal work.',
            status=LeaveApplication.Status.PENDING,
        )
        response = self.client.delete(f'/api/v1/admin/employees/{self.faculty.emp_id}/')
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertFalse(response.data['success'])
        self.assertIn('Set Inactive', response.data['message'])
        self.assertTrue(Employee.objects.filter(emp_id=self.faculty.emp_id).exists())

    def test_cannot_delete_own_employee_or_last_admin(self):
        response = self.client.delete(f'/api/v1/admin/employees/{self.admin.emp_id}/')
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('own employee', response.data['message'].lower())
        self.assertTrue(Employee.objects.filter(emp_id=self.admin.emp_id).exists())

    def test_assign_approvers_only_selected_employees(self):
        first = make_employee(
            emp_id='GCU010098',
            name='First Staff',
            designation='Administration',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='Establishment Section',
            email='first_adm@gcuniversity.ac.in',
            mobile='7002000098',
        )
        second = make_employee(
            emp_id='GCU010099',
            name='Second Staff',
            designation='Administration',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='Establishment Section',
            email='second_adm@gcuniversity.ac.in',
            mobile='7002000099',
        )
        approver_one = make_employee(
            emp_id='GCU010010',
            name='Staff Approver One',
            designation='Office Superintendent',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='',
            email='staff_approver1@gcuniversity.ac.in',
            mobile='7002000010',
        )
        approver_two = make_employee(
            emp_id='GCU010020',
            name='Staff Approver Two',
            designation='Deputy Registrar',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='',
            email='staff_approver2@gcuniversity.ac.in',
            mobile='7002000020',
        )
        dept_a1 = make_employee(
            emp_id='GCU010030',
            name='Department Approver One',
            designation='Section Officer',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='',
            email='dept_approver1@gcuniversity.ac.in',
            mobile='7002000030',
        )
        dept_a2 = make_employee(
            emp_id='GCU010031',
            name='Department Approver Two',
            designation='Section Officer',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='',
            email='dept_approver2@gcuniversity.ac.in',
            mobile='7002000031',
        )
        ApproverAssignment.objects.create(
            role=ApproverAssignment.Role.APPROVER_1,
            department='Administration',
            sub_department='',
            employee=dept_a1,
        )
        ApproverAssignment.objects.create(
            role=ApproverAssignment.Role.APPROVER_2,
            department='Administration',
            sub_department='',
            employee=dept_a2,
        )

        empty = self.client.post(
            '/api/v1/admin/employee-approvers/',
            {
                'emp_ids': [],
                'approver1_emp_id': approver_one.emp_id,
                'approver2_emp_id': approver_two.emp_id,
            },
            format='json',
        )
        self.assertEqual(empty.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertFalse(empty.data['success'])

        assigned = self.client.post(
            '/api/v1/admin/employee-approvers/',
            {
                'emp_ids': [first.emp_id],
                'approver1_emp_id': approver_one.emp_id,
                'approver2_emp_id': approver_two.emp_id,
            },
            format='json',
        )
        self.assertEqual(assigned.status_code, status.HTTP_200_OK)
        self.assertTrue(assigned.data['success'])
        self.assertEqual(assigned.data['data']['count'], 1)
        self.assertEqual(assigned.data['data']['emp_ids'], [first.emp_id])
        self.assertEqual(EmployeeApprover.objects.count(), 1)
        self.assertTrue(
            EmployeeApprover.objects.filter(
                employee=first,
                approver_1=approver_one,
                approver_2=approver_two,
            ).exists()
        )
        self.assertFalse(EmployeeApprover.objects.filter(employee=second).exists())
        self.assertEqual(find_approver1(first).emp_id, approver_one.emp_id)
        self.assertEqual(find_approver2(first).emp_id, approver_two.emp_id)
        self.assertEqual(find_approver1(second).emp_id, dept_a1.emp_id)
        self.assertEqual(find_approver2(second).emp_id, dept_a2.emp_id)

        listed = self.client.get(
            '/api/v1/admin/employee-approvers/',
            {'page_size': 500},
        )
        self.assertEqual(listed.status_code, status.HTTP_200_OK)
        self.assertGreaterEqual(listed.data['data']['count'], 2)
        listed_ids = [row['employee']['emp_id'] for row in listed.data['data']['results']]
        self.assertIn(first.emp_id, listed_ids)
        self.assertIn(second.emp_id, listed_ids)
        first_mapping = next(
            row for row in listed.data['data']['results'] if row['employee']['emp_id'] == first.emp_id
        )
        second_mapping = next(
            row for row in listed.data['data']['results'] if row['employee']['emp_id'] == second.emp_id
        )
        mapping_id = first_mapping['id']
        self.assertEqual(first_mapping['approver_1']['emp_id'], approver_one.emp_id)
        self.assertEqual(first_mapping['approver_2']['emp_id'], approver_two.emp_id)
        self.assertIsNone(second_mapping['id'])
        self.assertIsNone(second_mapping['approver_1'])
        self.assertIsNone(second_mapping['approver_2'])

        employees = self.client.get(
            '/api/v1/admin/employees/',
            {
                'department': 'Establishment Section',
                'q': 'First Staff',
                'status': 'Active',
            },
        )
        self.assertEqual(employees.status_code, status.HTTP_200_OK)
        emp_ids = [row['emp_id'] for row in employees.data['data']['results']]
        self.assertIn(first.emp_id, emp_ids)
        self.assertNotIn(second.emp_id, emp_ids)
        first_row = next(
            row for row in employees.data['data']['results'] if row['emp_id'] == first.emp_id
        )
        self.assertEqual(first_row['approver_1']['emp_id'], approver_one.emp_id)
        self.assertEqual(first_row['approver_2']['emp_id'], approver_two.emp_id)

        removed = self.client.delete(f'/api/v1/admin/employee-approvers/{mapping_id}/')
        self.assertEqual(removed.status_code, status.HTTP_200_OK)
        self.assertFalse(EmployeeApprover.objects.filter(employee=first).exists())
        self.assertEqual(find_approver1(first).emp_id, dept_a1.emp_id)

        bulk_alias = self.client.post(
            '/api/v1/admin/approvers/bulk/',
            {
                'emp_ids': [first.emp_id, second.emp_id],
                'approver1_emp_id': approver_one.emp_id,
                'approver2_emp_id': approver_two.emp_id,
            },
            format='json',
        )
        self.assertEqual(bulk_alias.status_code, status.HTTP_200_OK)
        self.assertEqual(EmployeeApprover.objects.count(), 2)

    def test_assign_personal_approver1_without_approver2(self):
        staff = make_employee(
            emp_id='GCU010198',
            name='Single Chain Staff',
            designation='Office Assistant',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='Establishment Section',
            email='single_chain_staff@gcuniversity.ac.in',
            mobile='7002000198',
        )
        approver_one = make_employee(
            emp_id='GCU010199',
            name='Only Approver One',
            designation='Office Superintendent',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='',
            email='only_approver_one@gcuniversity.ac.in',
            mobile='7002000199',
        )
        dept_a2 = make_employee(
            emp_id='GCU010200',
            name='Department Approver Two',
            designation='Deputy Registrar',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='',
            email='dept_a2_fallback@gcuniversity.ac.in',
            mobile='7002000200',
        )
        ApproverAssignment.objects.create(
            role=ApproverAssignment.Role.APPROVER_2,
            department='Administration',
            sub_department='',
            employee=dept_a2,
        )

        assigned = self.client.post(
            '/api/v1/admin/employee-approvers/',
            {
                'emp_ids': [staff.emp_id],
                'approver1_emp_id': approver_one.emp_id,
            },
            format='json',
        )
        self.assertEqual(assigned.status_code, status.HTTP_200_OK)
        self.assertTrue(assigned.data['success'])
        self.assertEqual(assigned.data['data']['count'], 1)
        mapping = EmployeeApprover.objects.get(employee=staff)
        self.assertEqual(mapping.approver_1_id, approver_one.id)
        self.assertIsNone(mapping.approver_2_id)
        self.assertEqual(find_approver1(staff).emp_id, approver_one.emp_id)
        self.assertIsNone(find_approver2(staff))

    def test_bulk_assign_lists_each_selected_employee(self):
        first = make_employee(
            emp_id='GCU010401',
            name='Bulk Staff One',
            designation='Office Assistant',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='Establishment Section',
            email='bulk_staff_one@gcuniversity.ac.in',
            mobile='7002000401',
        )
        second = make_employee(
            emp_id='GCU010402',
            name='Bulk Staff Two',
            designation='Office Assistant',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='Establishment Section',
            email='bulk_staff_two@gcuniversity.ac.in',
            mobile='7002000402',
        )
        third = make_employee(
            emp_id='GCU010403',
            name='Bulk Staff Three',
            designation='Office Assistant',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='Establishment Section',
            email='bulk_staff_three@gcuniversity.ac.in',
            mobile='7002000403',
        )
        approver_one = make_employee(
            emp_id='GCU010404',
            name='Bulk Approver One',
            designation='Office Superintendent',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='',
            email='bulk_approver_one@gcuniversity.ac.in',
            mobile='7002000404',
        )
        approver_two = make_employee(
            emp_id='GCU010405',
            name='Bulk Approver Two',
            designation='Deputy Registrar',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='',
            email='bulk_approver_two@gcuniversity.ac.in',
            mobile='7002000405',
        )

        assigned = self.client.post(
            '/api/v1/admin/employee-approvers/',
            {
                'emp_ids': [first.emp_id, second.emp_id],
                'approver1_emp_id': approver_one.emp_id,
                'approver2_emp_id': approver_two.emp_id,
            },
            format='json',
        )
        self.assertEqual(assigned.status_code, status.HTTP_200_OK)
        self.assertTrue(assigned.data['success'])
        self.assertEqual(assigned.data['data']['count'], 2)
        self.assertEqual(
            set(assigned.data['data']['emp_ids']),
            {first.emp_id, second.emp_id},
        )
        self.assertEqual(EmployeeApprover.objects.count(), 2)
        self.assertTrue(
            EmployeeApprover.objects.filter(
                employee=first,
                approver_1=approver_one,
                approver_2=approver_two,
            ).exists()
        )
        self.assertTrue(
            EmployeeApprover.objects.filter(
                employee=second,
                approver_1=approver_one,
                approver_2=approver_two,
            ).exists()
        )
        self.assertFalse(EmployeeApprover.objects.filter(employee=third).exists())
        self.assertEqual(find_approver1(first).emp_id, approver_one.emp_id)
        self.assertEqual(find_approver2(second).emp_id, approver_two.emp_id)
        self.assertIsNone(find_approver1(third))
        self.assertIsNone(find_approver2(third))

        listed = self.client.get(
            '/api/v1/admin/employee-approvers/',
            {'department': 'Establishment Section', 'page_size': 500},
        )
        self.assertEqual(listed.status_code, status.HTTP_200_OK)
        by_id = {
            row['employee']['emp_id']: row for row in listed.data['data']['results']
        }
        self.assertIn(first.emp_id, by_id)
        self.assertIn(second.emp_id, by_id)
        self.assertIn(third.emp_id, by_id)
        self.assertEqual(by_id[first.emp_id]['approver_1']['emp_id'], approver_one.emp_id)
        self.assertEqual(by_id[first.emp_id]['approver_2']['emp_id'], approver_two.emp_id)
        self.assertEqual(by_id[second.emp_id]['approver_1']['emp_id'], approver_one.emp_id)
        self.assertEqual(by_id[second.emp_id]['approver_2']['emp_id'], approver_two.emp_id)
        self.assertTrue(by_id[first.emp_id]['id'])
        self.assertTrue(by_id[second.emp_id]['id'])
        self.assertNotEqual(by_id[first.emp_id]['id'], by_id[second.emp_id]['id'])
        self.assertIsNone(by_id[third.emp_id]['id'])
        self.assertIsNone(by_id[third.emp_id]['approver_1'])
        self.assertIsNone(by_id[third.emp_id]['approver_2'])

        listed_ids = [row['employee']['emp_id'] for row in listed.data['data']['results']]
        self.assertLess(listed_ids.index(first.emp_id), listed_ids.index(third.emp_id))
        self.assertLess(listed_ids.index(second.emp_id), listed_ids.index(third.emp_id))

    def test_employee_approver_filters_and_isolated_delete(self):
        first = make_employee(
            emp_id='GCU010198',
            name='Mapped Alpha',
            designation='Office Assistant',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='Establishment Section',
            email='mapped_alpha@gcuniversity.ac.in',
            mobile='7002000198',
        )
        second = make_employee(
            emp_id='GCU030198',
            name='Mapped Beta',
            designation='Laboratory Assistant',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='School Of Engineering and Technology',
            sub_department='Department of Mechanical Engineering',
            email='mapped_beta@gcuniversity.ac.in',
            mobile='7002000199',
        )
        approver_one = make_employee(
            emp_id='GCU010110',
            name='Isolated Approver One',
            designation='Office Superintendent',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='',
            email='isolated_a1@gcuniversity.ac.in',
            mobile='7002000110',
        )
        approver_two = make_employee(
            emp_id='GCU010120',
            name='Isolated Approver Two',
            designation='Deputy Registrar',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='',
            email='isolated_a2@gcuniversity.ac.in',
            mobile='7002000120',
        )
        other_a1 = make_employee(
            emp_id='GCU030110',
            name='Other Approver One',
            designation='Section Officer',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='School Of Engineering and Technology',
            sub_department='',
            email='other_a1@gcuniversity.ac.in',
            mobile='7002000310',
        )
        other_a2 = make_employee(
            emp_id='GCU030120',
            name='Other Approver Two',
            designation='Section Officer',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='School Of Engineering and Technology',
            sub_department='',
            email='other_a2@gcuniversity.ac.in',
            mobile='7002000320',
        )
        EmployeeApprover.objects.create(
            employee=first,
            approver_1=approver_one,
            approver_2=approver_two,
        )
        EmployeeApprover.objects.create(
            employee=second,
            approver_1=other_a1,
            approver_2=other_a2,
        )

        listed = self.client.get(
            '/api/v1/admin/employee-approvers/',
            {'page_size': 500},
        )
        self.assertEqual(listed.status_code, status.HTTP_200_OK)
        self.assertTrue(listed.data['success'])
        self.assertIn('results', listed.data['data'])
        self.assertGreaterEqual(listed.data['data']['count'], 2)
        listed_ids = [row['employee']['emp_id'] for row in listed.data['data']['results']]
        self.assertIn(first.emp_id, listed_ids)
        self.assertIn(second.emp_id, listed_ids)

        by_name = self.client.get('/api/v1/admin/employee-approvers/', {'q': 'Mapped Alpha'})
        self.assertEqual(by_name.data['data']['count'], 1)
        self.assertEqual(by_name.data['data']['results'][0]['employee']['emp_id'], first.emp_id)

        by_id = self.client.get('/api/v1/admin/employee-approvers/', {'q': first.emp_id})
        self.assertEqual(by_id.data['data']['count'], 1)
        self.assertEqual(by_id.data['data']['results'][0]['employee']['emp_id'], first.emp_id)

        by_dept = self.client.get(
            '/api/v1/admin/employee-approvers/',
            {'department': 'Establishment Section', 'page_size': 500},
        )
        dept_ids = [row['employee']['emp_id'] for row in by_dept.data['data']['results']]
        self.assertIn(first.emp_id, dept_ids)
        self.assertNotIn(second.emp_id, dept_ids)

        by_sub = self.client.get(
            '/api/v1/admin/employee-approvers/',
            {'sub_department': 'Establishment Section', 'page_size': 500},
        )
        self.assertIn(
            first.emp_id,
            [row['employee']['emp_id'] for row in by_sub.data['data']['results']],
        )

        by_desig = self.client.get(
            '/api/v1/admin/employee-approvers/',
            {'designation': 'Laboratory Assistant'},
        )
        self.assertEqual(by_desig.data['data']['count'], 1)
        self.assertEqual(by_desig.data['data']['results'][0]['employee']['emp_id'], second.emp_id)

        combined = self.client.get(
            '/api/v1/admin/employee-approvers/',
            {
                'q': 'Mapped',
                'department': 'Establishment Section',
                'designation': 'Laboratory Assistant',
            },
        )
        self.assertEqual(combined.data['data']['count'], 0)

        first_id = by_name.data['data']['results'][0]['id']
        second_id = by_desig.data['data']['results'][0]['id']

        cleared = self.client.patch(
            f'/api/v1/admin/employee-approvers/{first_id}/',
            {'clear_approver_1': True},
            format='json',
        )
        self.assertEqual(cleared.status_code, status.HTTP_200_OK)
        self.assertIsNone(cleared.data['data']['approver_1'])
        self.assertEqual(cleared.data['data']['approver_2']['emp_id'], approver_two.emp_id)
        self.assertEqual(cleared.data['data']['employee']['emp_id'], first.emp_id)
        first_row = EmployeeApprover.objects.get(pk=first_id)
        self.assertIsNone(first_row.approver_1_id)
        self.assertEqual(first_row.approver_2_id, approver_two.id)
        second_row = EmployeeApprover.objects.get(pk=second_id)
        self.assertEqual(second_row.approver_1_id, other_a1.id)
        self.assertEqual(second_row.approver_2_id, other_a2.id)

        removed = self.client.delete(f'/api/v1/admin/employee-approvers/{first_id}/')
        self.assertEqual(removed.status_code, status.HTTP_200_OK)
        self.assertFalse(EmployeeApprover.objects.filter(pk=first_id).exists())
        self.assertTrue(
            EmployeeApprover.objects.filter(
                pk=second_id,
                employee=second,
                approver_1=other_a1,
                approver_2=other_a2,
            ).exists()
        )
        self.assertEqual(EmployeeApprover.objects.count(), 1)

        still_listed = self.client.get(
            '/api/v1/admin/employee-approvers/',
            {'q': first.emp_id},
        )
        self.assertEqual(still_listed.data['data']['count'], 1)
        self.assertIsNone(still_listed.data['data']['results'][0]['approver_1'])
        self.assertIsNone(still_listed.data['data']['results'][0]['approver_2'])

        both_gone = self.client.patch(
            f'/api/v1/admin/employee-approvers/{second_id}/',
            {'clear_approver_1': True, 'clear_approver_2': True},
            format='json',
        )
        self.assertEqual(both_gone.status_code, status.HTTP_200_OK)
        self.assertIsNone(both_gone.data['data']['id'])
        self.assertIsNone(both_gone.data['data']['approver_1'])
        self.assertIsNone(both_gone.data['data']['approver_2'])
        self.assertEqual(both_gone.data['data']['employee']['emp_id'], second.emp_id)
        self.assertFalse(EmployeeApprover.objects.filter(pk=second_id).exists())
        self.assertEqual(EmployeeApprover.objects.count(), 0)

    def test_employee_approver_export_uses_filters(self):
        mapped = make_employee(
            emp_id='GCU010198',
            name='Mapped Alpha',
            designation='Office Assistant',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='Establishment Section',
            email='mapped_alpha@gcuniversity.ac.in',
            mobile='7002000198',
        )
        hod = make_employee(
            emp_id='GCU020198',
            name='Mapped HOD',
            designation='Professor & HOD',
            designation_type=Employee.DesignationType.FACULTY,
            group_name='Faculty',
            department='School Of Engineering and Technology',
            sub_department='Department of Mechanical Engineering',
            email='mapped_hod@gcuniversity.ac.in',
            mobile='7002000298',
        )
        approver_one = make_employee(
            emp_id='GCU010110',
            name='Export Approver One',
            designation='Office Superintendent',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='',
            email='export_a1@gcuniversity.ac.in',
            mobile='7002000110',
        )
        approver_two = make_employee(
            emp_id='GCU010120',
            name='Export Approver Two',
            designation='Deputy Registrar',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='',
            email='export_a2@gcuniversity.ac.in',
            mobile='7002000120',
        )
        EmployeeApprover.objects.create(
            employee=mapped,
            approver_1=approver_one,
            approver_2=approver_two,
        )

        exported = self.client.get(
            '/api/v1/admin/employee-approvers/export/',
            {'q': 'Mapped Alpha'},
        )
        self.assertEqual(exported.status_code, status.HTTP_200_OK)
        self.assertIn(
            'employee-approvers.xlsx',
            exported['Content-Disposition'],
        )
        workbook = load_workbook(BytesIO(exported.content))
        sheet = workbook.active
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        self.assertEqual(
            headers,
            [
                'Employee ID',
                'Name',
                'Department',
                'Designation',
                'Approver 1',
                'Approver 1 ID',
                'Approver 2',
                'Approver 2 ID',
            ],
        )
        self.assertEqual(sheet.max_row, 2)
        values = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
        self.assertEqual(values[0], mapped.emp_id)
        self.assertEqual(values[1], mapped.name)
        self.assertEqual(values[2], mapped.sub_department)
        self.assertEqual(values[3], mapped.designation)
        self.assertEqual(values[4], approver_one.name)
        self.assertEqual(values[5], approver_one.emp_id)
        self.assertEqual(values[6], approver_two.name)
        self.assertEqual(values[7], approver_two.emp_id)

        hod_export = self.client.get(
            '/api/v1/admin/employee-approvers/export/',
            {'designation': 'Professor & HOD'},
        )
        hod_row = [cell.value for cell in next(
            load_workbook(BytesIO(hod_export.content)).active.iter_rows(min_row=2, max_row=2)
        )]
        self.assertEqual(hod_row[0], hod.emp_id)
        self.assertEqual(hod_row[4], 'Vice-Chancellor only')
        self.assertFalse(hod_row[5])
        self.assertFalse(hod_row[6])

        login_as(self.client, 'GCU020041')
        denied = self.client.get('/api/v1/admin/employee-approvers/export/')
        self.assertEqual(denied.status_code, status.HTTP_403_FORBIDDEN)
        login_as(self.client, 'GCU090001', 'NewPass#2026')

    def test_admin_reset_password(self):
        target = make_employee(
            emp_id='GCU099901',
            name='Password Reset Dummy',
            designation='Office Assistant',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='',
            email='reset_dummy@gcuniversity.ac.in',
            mobile='7002999901',
        )
        new_password = 'TempReset#2026'

        login_as(self.client, 'GCU020041')
        denied = self.client.post(
            f'/api/v1/admin/employees/{target.emp_id}/reset-password/',
            {'new_password': new_password, 'confirm_password': new_password},
            format='json',
        )
        self.assertEqual(denied.status_code, status.HTTP_403_FORBIDDEN)

        login_as(self.client, 'GCU090001', 'NewPass#2026')
        reset = self.client.post(
            f'/api/v1/admin/employees/{target.emp_id}/reset-password/',
            {'new_password': new_password, 'confirm_password': new_password},
            format='json',
        )
        self.assertEqual(reset.status_code, status.HTTP_200_OK)
        self.assertTrue(reset.data['success'])
        self.assertTrue(reset.data['data']['must_change_password'])
        self.assertEqual(reset.data['data']['emp_id'], target.emp_id)
        self.assertNotIn('new_password', reset.data)
        self.assertNotIn('new_password', reset.data.get('data') or {})

        mismatch = self.client.post(
            f'/api/v1/admin/employees/{target.emp_id}/reset-password/',
            {'new_password': new_password, 'confirm_password': 'OtherPass#2026'},
            format='json',
        )
        self.assertEqual(mismatch.status_code, status.HTTP_400_BAD_REQUEST)

        self.client.credentials()
        login = self.client.post(
            '/api/v1/auth/login/',
            {'emp_id': target.emp_id, 'password': new_password},
            format='json',
        )
        self.assertEqual(login.status_code, status.HTTP_200_OK)
        self.assertTrue(login.data['data']['user']['must_change_password'])
        self.assertTrue(Account.objects.get(employee=target).must_change_password)


LEAVE_EXCEL_HEADERS = [
    'Serial No.',
    'Employee ID',
    'Name',
    'Location',
    'Leave Type',
    'Request Date',
    'From Date',
    'To Date',
    'Status',
    'Timeline',
    'Total Days',
]


class LoginPageCopyTests(SimpleTestCase):
    def test_login_page_omits_default_password_hint(self):
        path = Path(settings.BASE_DIR).parent / 'frontend' / 'src' / 'pages' / 'auth' / 'Login.jsx'
        text = path.read_text(encoding='utf-8')
        self.assertNotIn('gcu@123', text)


class AdminLeaveExportTests(APITestCase):
    def setUp(self):
        self.faculty = make_employee()
        self.admin = make_employee(
            emp_id='GCU090001',
            name='Vice Chancellor',
            designation='Vice-Chancellor',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='',
            email='vc@gcuniversity.ac.in',
            mobile='9000000001',
        )
        login_as(self.client, 'GCU090001')
        account = Account.objects.get(employee=self.admin)
        account.is_admin = True
        account.save(update_fields=['is_admin', 'updated_at'])
        account.user.is_staff = True
        account.user.save(update_fields=['is_staff'])
        login_as(self.client, 'GCU090001', 'NewPass#2026')
        self.casual, _ = LeaveType.objects.get_or_create(
            code='CL',
            defaults={
                'name': 'Casual Leave',
                'max_days_per_year': 12,
                'applicable_to': LeaveType.ApplicableTo.ALL,
                'sort_order': 1,
            },
        )
        self.application = LeaveApplication.objects.create(
            employee=self.faculty,
            leave_type=self.casual,
            start_date=date(2026, 3, 10),
            end_date=date(2026, 3, 11),
            days=2,
            reason='Conference travel.',
            status=LeaveApplication.Status.APPROVED,
            requested_on=date(2026, 3, 1),
        )
        ApprovalAction.objects.create(
            application=self.application,
            step=ApprovalAction.Step.HOD,
            actor=self.admin,
            decision=ApprovalAction.Decision.APPROVED,
        )
        self.other = make_employee(
            emp_id='GCU010077',
            name='Other Person',
            designation='Clerk',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='Establishment Section',
            academy='Guwahati Campus',
            email='other_adm@gcuniversity.ac.in',
            mobile='7002000077',
        )
        LeaveApplication.objects.create(
            employee=self.other,
            leave_type=self.casual,
            start_date=date(2026, 6, 1),
            end_date=date(2026, 6, 8),
            days=8,
            reason='Family matter.',
            status=LeaveApplication.Status.PENDING,
            requested_on=date(2026, 5, 20),
        )

    def test_non_admin_cannot_list_or_export_leaves(self):
        login_as(self.client, 'GCU020041')
        listed = self.client.get('/api/v1/admin/leaves/')
        exported = self.client.get('/api/v1/admin/leaves/export/')
        self.assertEqual(listed.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(exported.status_code, status.HTTP_403_FORBIDDEN)

    def test_admin_leaves_list_and_export_columns(self):
        listed = self.client.get('/api/v1/admin/leaves/', {'q': 'Bhabajit'})
        self.assertEqual(listed.status_code, status.HTTP_200_OK)
        self.assertTrue(listed.data['success'])
        results = listed.data['data']['results']
        self.assertEqual(len(results), 1)
        row = results[0]
        self.assertEqual(row['serial_no'], 1)
        self.assertEqual(row['emp_id'], 'GCU020041')
        self.assertEqual(row['name'], 'Bhabajit Baruah')
        self.assertEqual(row['location'], 'Girijananda Chowdhury University-Assam')
        self.assertEqual(row['leave_type'], 'Casual Leave')
        self.assertEqual(str(row['request_date']), '2026-03-01')
        self.assertEqual(str(row['from_date']), '2026-03-10')
        self.assertEqual(str(row['to_date']), '2026-03-11')
        self.assertEqual(row['status'], 'APPROVED')
        self.assertEqual(row['status_label'], 'Approved')
        self.assertEqual(row['timeline'], 'View')
        self.assertEqual(row['total_days'], 2)
        self.assertEqual(row['actions'][0]['step'], 'HOD')
        self.assertEqual(row['actions'][0]['decision_label'], 'Approved')
        for key in (
            'serial_no', 'emp_id', 'name', 'location', 'leave_type',
            'request_date', 'from_date', 'to_date', 'status', 'timeline', 'total_days',
        ):
            self.assertIn(key, row)

        by_dept = self.client.get(
            '/api/v1/admin/leaves/',
            {'sub_department': self.faculty.sub_department},
        )
        self.assertEqual(len(by_dept.data['data']['results']), 1)
        self.assertEqual(by_dept.data['data']['results'][0]['emp_id'], 'GCU020041')

        by_desig = self.client.get('/api/v1/admin/leaves/', {'designation': 'Clerk'})
        self.assertEqual(len(by_desig.data['data']['results']), 1)
        self.assertEqual(by_desig.data['data']['results'][0]['location'], 'Guwahati Campus')

        overlapping = self.client.get(
            '/api/v1/admin/leaves/',
            {'start': '2026-03-11', 'end': '2026-03-20'},
        )
        self.assertEqual(
            [item['emp_id'] for item in overlapping.data['data']['results']],
            ['GCU020041'],
        )
        start_only = self.client.get('/api/v1/admin/leaves/', {'start': '2026-06-01'})
        self.assertEqual(
            [item['emp_id'] for item in start_only.data['data']['results']],
            ['GCU010077'],
        )

        exported = self.client.get('/api/v1/admin/leaves/export/', {'q': 'Bhabajit'})
        self.assertEqual(exported.status_code, status.HTTP_200_OK)
        self.assertIn(
            'leave-applications.xlsx',
            exported['Content-Disposition'],
        )
        workbook = load_workbook(BytesIO(exported.content))
        sheet = workbook.active
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        self.assertEqual(headers, LEAVE_EXCEL_HEADERS)
        values = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
        self.assertEqual(values[0], 1)
        self.assertEqual(values[1], 'GCU020041')
        self.assertEqual(values[2], 'Bhabajit Baruah')
        self.assertEqual(values[3], 'Girijananda Chowdhury University-Assam')
        self.assertEqual(values[4], 'Casual Leave')
        self.assertEqual(values[5], '01-03-2026')
        self.assertEqual(values[6], '10-03-2026')
        self.assertEqual(values[7], '11-03-2026')
        self.assertEqual(values[8], 'Approved')
        self.assertEqual(values[9], 'View')
        self.assertEqual(values[10], 2)
        self.assertEqual(sheet.max_row, 2)


class AccessRolesAndAssignLeaveTests(APITestCase):
    def setUp(self):
        self.faculty = make_employee()
        self.operator_emp = make_employee(
            emp_id='GCU020099',
            name='Operator One',
            designation='Office Assistant',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='Registrar Office',
            email='operator@gcuniversity.ac.in',
            mobile='9000000099',
        )
        self.hod = make_employee(
            emp_id='GCU020014',
            name='Debarshi Mallick',
            designation='Associate Professor & HOD',
            email='debarshi_me@gcuniversity.ac.in',
            mobile='7002474854',
        )
        self.admin = make_employee(
            emp_id='GCU090001',
            name='Vice Chancellor',
            designation='Vice-Chancellor',
            designation_type=Employee.DesignationType.STAFF,
            group_name='Admin',
            department='Administration',
            sub_department='',
            email='vc@gcuniversity.ac.in',
            mobile='9000000001',
        )
        ApproverAssignment.objects.create(
            role=ApproverAssignment.Role.HOD,
            department=self.faculty.department,
            sub_department=self.faculty.sub_department,
            employee=self.hod,
        )
        ApproverAssignment.objects.create(
            role=ApproverAssignment.Role.VC,
            department='',
            sub_department='',
            employee=self.admin,
        )
        self.casual, _ = LeaveType.objects.get_or_create(
            code='CL',
            defaults={
                'name': 'Casual Leave',
                'max_days_per_year': 12,
                'applicable_to': LeaveType.ApplicableTo.ALL,
                'sort_order': 1,
            },
        )
        login_as(self.client, 'GCU090001')
        admin_account = Account.objects.get(employee=self.admin)
        admin_account.is_admin = True
        admin_account.save(update_fields=['is_admin', 'updated_at'])
        admin_account.user.is_staff = True
        admin_account.user.save(update_fields=['is_staff'])
        login_as(self.client, 'GCU090001', 'NewPass#2026')
        granted = self.client.patch(
            '/api/v1/admin/access/GCU020099/',
            {'is_operator': True},
            format='json',
        )
        self.assertEqual(granted.status_code, status.HTTP_200_OK)
        self.assertTrue(granted.data['data']['is_operator'])
        self.assertFalse(granted.data['data']['is_admin'])

    def test_login_payload_includes_is_operator(self):
        login_as(self.client, 'GCU020099')
        me = self.client.get('/api/v1/auth/me/')
        self.assertEqual(me.status_code, status.HTTP_200_OK)
        self.assertTrue(me.data['data']['is_operator'])
        self.assertFalse(me.data['data']['is_admin'])

        login_as(self.client, 'GCU020041')
        faculty_me = self.client.get('/api/v1/auth/me/')
        self.assertFalse(faculty_me.data['data']['is_operator'])
        self.assertFalse(faculty_me.data['data']['is_admin'])

    def test_faculty_cannot_access_admin_or_assign_leave(self):
        login_as(self.client, 'GCU020041')
        overview = self.client.get('/api/v1/admin/overview/')
        self.assertEqual(overview.status_code, status.HTTP_403_FORBIDDEN)
        assigned = self.client.post(
            '/api/v1/admin/assign-leave/',
            {
                'emp_id': 'GCU020041',
                'leave_type_id': self.casual.id,
                'days': 1,
                'action': 'CREDIT',
            },
            format='json',
        )
        self.assertEqual(assigned.status_code, status.HTTP_403_FORBIDDEN)

    def test_operator_can_use_analytics_but_not_settings(self):
        login_as(self.client, 'GCU020099')
        overview = self.client.get('/api/v1/admin/overview/')
        self.assertEqual(overview.status_code, status.HTTP_200_OK)
        leaves = self.client.get('/api/v1/admin/leaves/')
        self.assertEqual(leaves.status_code, status.HTTP_200_OK)
        listed = self.client.get('/api/v1/admin/employees/')
        self.assertEqual(listed.status_code, status.HTTP_403_FORBIDDEN)
        denied = self.client.patch(
            '/api/v1/admin/access/GCU020041/',
            {'is_admin': True},
            format='json',
        )
        self.assertEqual(denied.status_code, status.HTTP_403_FORBIDDEN)
        employee_patch = self.client.patch(
            '/api/v1/admin/employees/GCU020041/',
            {'is_admin': True, 'name': 'Bhabajit Baruah'},
            format='json',
        )
        self.assertEqual(employee_patch.status_code, status.HTTP_403_FORBIDDEN)
        assigned = self.client.post(
            '/api/v1/admin/assign-leave/',
            {
                'emp_id': 'GCU020041',
                'leave_type_id': self.casual.id,
                'days': 1,
                'action': 'CREDIT',
            },
            format='json',
        )
        self.assertEqual(assigned.status_code, status.HTTP_403_FORBIDDEN)

    def test_approver_without_operator_cannot_use_settings(self):
        login_as(self.client, 'GCU020014')
        denied = self.client.get('/api/v1/admin/overview/')
        self.assertEqual(denied.status_code, status.HTTP_403_FORBIDDEN)
        inbox = self.client.get('/api/v1/approvals/')
        self.assertEqual(inbox.status_code, status.HTTP_200_OK)

    def test_credit_raises_yearly_cap(self):
        login_as(self.client, 'GCU020041')
        start = date.today() + timedelta(days=40)
        end = start + timedelta(days=12)
        too_many = self.client.post(
            '/api/v1/leaves/',
            {
                'leave_type_id': self.casual.id,
                'start_date': start.isoformat(),
                'end_date': end.isoformat(),
                'reason': 'Extended family function in Guwahati.',
            },
            format='json',
        )
        self.assertEqual(too_many.status_code, status.HTTP_400_BAD_REQUEST)

        login_as(self.client, 'GCU090001', 'NewPass#2026')
        credited = self.client.post(
            '/api/v1/admin/assign-leave/',
            {
                'emp_id': 'GCU020041',
                'leave_type_id': self.casual.id,
                'days': 5,
                'action': 'CREDIT',
            },
            format='json',
        )
        self.assertEqual(credited.status_code, status.HTTP_201_CREATED)
        cl_balance = next(
            item for item in credited.data['data']['balances'] if item['code'] == 'CL'
        )
        self.assertEqual(cl_balance['max_days_per_year'], 17)
        self.assertEqual(cl_balance['remaining'], 17)

        login_as(self.client, 'GCU020041')
        allowed = self.client.post(
            '/api/v1/leaves/',
            {
                'leave_type_id': self.casual.id,
                'start_date': start.isoformat(),
                'end_date': end.isoformat(),
                'reason': 'Extended family function in Guwahati.',
            },
            format='json',
        )
        self.assertEqual(allowed.status_code, status.HTTP_201_CREATED)
        summary = self.client.get('/api/v1/leaves/summary/')
        cl_after = next(
            item for item in summary.data['data']['balances'] if item['code'] == 'CL'
        )
        self.assertEqual(cl_after['used'], 13)
        self.assertEqual(cl_after['remaining'], 4)

    def test_sanction_creates_approved_leave(self):
        login_as(self.client, 'GCU090001', 'NewPass#2026')
        start = date.today() + timedelta(days=50)
        sanctioned = self.client.post(
            '/api/v1/admin/assign-leave/',
            {
                'emp_id': 'GCU020041',
                'leave_type_id': self.casual.id,
                'days': 2,
                'action': 'SANCTION',
                'start_date': start.isoformat(),
            },
            format='json',
        )
        self.assertEqual(sanctioned.status_code, status.HTTP_201_CREATED)
        data = sanctioned.data['data']
        self.assertEqual(data['kind'], 'SANCTION')
        self.assertEqual(data['status'], LeaveApplication.Status.APPROVED)
        self.assertEqual(data['waiting_on'], None)
        self.assertEqual(data['source'], LeaveApplication.Source.ASSIGNED)
        self.assertEqual(data['days'], 2)
        application = LeaveApplication.objects.get(pk=data['id'])
        self.assertEqual(application.status, LeaveApplication.Status.APPROVED)
        self.assertEqual(application.waiting_on_id, None)
        self.assertEqual(application.current_step, '')
        self.assertTrue(
            Notification.objects.filter(employee=self.faculty, title__icontains='sanctioned').exists()
        )
        cl_balance = next(
            item for item in data['balances'] if item['code'] == 'CL'
        )
        self.assertEqual(cl_balance['used'], 2)
        self.assertEqual(cl_balance['remaining'], 10)

    def test_credit_on_uncapped_type_shows_dashboard_balance(self):
        special, _ = LeaveType.objects.get_or_create(
            code='SPL',
            defaults={
                'name': 'Special Leave',
                'max_days_per_year': None,
                'applicable_to': LeaveType.ApplicableTo.ALL,
                'sort_order': 11,
            },
        )
        special.max_days_per_year = None
        special.save(update_fields=['max_days_per_year'])
        login_as(self.client, 'GCU090001', 'NewPass#2026')
        credited = self.client.post(
            '/api/v1/admin/assign-leave/',
            {
                'emp_id': 'GCU020041',
                'leave_type_id': special.id,
                'days': 3,
                'action': 'CREDIT',
            },
            format='json',
        )
        self.assertEqual(credited.status_code, status.HTTP_201_CREATED)
        spl = next(item for item in credited.data['data']['balances'] if item['code'] == 'SPL')
        self.assertEqual(spl['max_days_per_year'], 3)
        self.assertEqual(spl['remaining'], 3)
        self.assertEqual(spl['used'], 0)

        login_as(self.client, 'GCU020041')
        summary = self.client.get('/api/v1/leaves/summary/')
        spl_after = next(
            item for item in summary.data['data']['balances'] if item['code'] == 'SPL'
        )
        self.assertEqual(spl_after['remaining'], 3)
        self.assertEqual(spl_after['max_days_per_year'], 3)
        self.assertEqual(spl_after['used'], 0)

    def test_credit_on_uncapped_type_stays_available_after_existing_used(self):
        duty, _ = LeaveType.objects.get_or_create(
            code='DL',
            defaults={
                'name': 'Duty Leave',
                'max_days_per_year': None,
                'applicable_to': LeaveType.ApplicableTo.ALL,
                'sort_order': 4,
            },
        )
        duty.max_days_per_year = None
        duty.save(update_fields=['max_days_per_year'])
        start = date.today()
        LeaveApplication.objects.create(
            employee=self.faculty,
            leave_type=duty,
            start_date=start,
            end_date=start + timedelta(days=1),
            days=2,
            reason='Official university duty.',
            status=LeaveApplication.Status.APPROVED,
            source=LeaveApplication.Source.INTERNAL,
        )
        login_as(self.client, 'GCU090001', 'NewPass#2026')
        credited = self.client.post(
            '/api/v1/admin/assign-leave/',
            {
                'emp_id': 'GCU020041',
                'leave_type_id': duty.id,
                'days': 1,
                'action': 'CREDIT',
            },
            format='json',
        )
        self.assertEqual(credited.status_code, status.HTTP_201_CREATED)
        dl = next(item for item in credited.data['data']['balances'] if item['code'] == 'DL')
        self.assertEqual(dl['used'], 2)
        self.assertEqual(dl['remaining'], 1)
        self.assertEqual(dl['max_days_per_year'], 1)

    def test_credited_uncapped_leave_cannot_exceed_remaining(self):
        special, _ = LeaveType.objects.get_or_create(
            code='SPL',
            defaults={
                'name': 'Special Leave',
                'max_days_per_year': None,
                'applicable_to': LeaveType.ApplicableTo.ALL,
                'sort_order': 11,
            },
        )
        special.max_days_per_year = None
        special.save(update_fields=['max_days_per_year'])
        login_as(self.client, 'GCU090001', 'NewPass#2026')
        credited = self.client.post(
            '/api/v1/admin/assign-leave/',
            {
                'emp_id': 'GCU020041',
                'leave_type_id': special.id,
                'days': 14,
                'action': 'CREDIT',
            },
            format='json',
        )
        self.assertEqual(credited.status_code, status.HTTP_201_CREATED)

        login_as(self.client, 'GCU020041')
        types = self.client.get('/api/v1/leave-types/')
        spl_type = next(item for item in types.data['data'] if item['code'] == 'SPL')
        self.assertEqual(spl_type['max_days_per_year'], 14)
        self.assertEqual(spl_type['remaining'], 14)

        start = date.today() + timedelta(days=70)
        too_many = self.client.post(
            '/api/v1/leaves/',
            {
                'leave_type_id': special.id,
                'start_date': start.isoformat(),
                'end_date': (start + timedelta(days=16)).isoformat(),
                'reason': 'Personal work requiring more days than credited.',
            },
            format='json',
        )
        self.assertEqual(too_many.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('remaining', too_many.data['message'].lower())

        allowed = self.client.post(
            '/api/v1/leaves/',
            {
                'leave_type_id': special.id,
                'start_date': start.isoformat(),
                'end_date': (start + timedelta(days=13)).isoformat(),
                'reason': 'Personal work within the credited special leave.',
            },
            format='json',
        )
        self.assertEqual(allowed.status_code, status.HTTP_201_CREATED)
        self.assertEqual(allowed.data['data']['days'], 14)

